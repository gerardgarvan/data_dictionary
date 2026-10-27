"""
Tests for python/pdd/dictionary_workbook.py.

All fixtures use 13-column profile_PCNR.csv schema:
  family, variable, var_class, stat, level, n, value,
  pct_total, pct_nonmiss, sort_order, small_cell_flag, date_kind, display_format

n_nonmiss and n_distinct are DERIVED outputs, not CSV columns.
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pandas as pd
import pytest

from pdd.dictionary_workbook import (
    DEFAULT_DISPLAY_CONFIG,
    SHEET_ORDER,
    _key_rows,
    apply_age_topcode,
    apply_suppression,
    build_qc_sheet,
    build_sources_sheet,
    build_variables_sheet,
    write_workbook,
)

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

PROFILE_COLS = [
    "family", "variable", "var_class", "stat", "level",
    "n", "value", "pct_total", "pct_nonmiss", "sort_order",
    "small_cell_flag", "date_kind", "display_format",
]


def _make_rows(*rows):
    """Return DataFrame with profile columns from list-of-dicts (fills missing with '')."""
    recs = []
    for r in rows:
        rec = {c: "" for c in PROFILE_COLS}
        rec.update(r)
        recs.append(rec)
    return pd.DataFrame(recs, columns=PROFILE_COLS)


def _make_age_csv(tmp_path, variables):
    """Write a minimal age_variables.csv to tmp_path; return the directory."""
    p = tmp_path / "age_variables.csv"
    p.write_text("variable,reason\n" + "\n".join(f"{v},test" for v in variables))
    return tmp_path


def _cfg(**kwargs):
    cfg = dict(DEFAULT_DISPLAY_CONFIG)
    cfg.update(kwargs)
    return cfg


# ---------------------------------------------------------------------------
# Age file hard fail
# ---------------------------------------------------------------------------

def test_age_topcode_missing_file_raises(tmp_path):
    """apply_age_topcode with a nonexistent age file must raise (not silently return empty)."""
    missing = tmp_path / "no_such_dir" / "age_variables.csv"
    df = _make_rows({"variable": "PCNR_AGE", "var_class": "CONT_NUM", "stat": "LEVEL",
                     "level": "91", "n": "5", "small_cell_flag": "Y"})
    with pytest.raises((FileNotFoundError, ValueError)) as exc_info:
        apply_age_topcode(df, missing)
    assert str(missing) in str(exc_info.value)


# ---------------------------------------------------------------------------
# Transform order: topcode BEFORE suppression
# ---------------------------------------------------------------------------

def test_transform_order_topcode_before_suppress(tmp_path):
    """
    Age var with LEVEL rows 90, 91, 92 each n=5 (individually small-cell).
    After topcode: merged '90+' row has n=15 (not small-cell).
    After suppression: '90+' row is NOT suppressed.
    """
    docs_dir = _make_age_csv(tmp_path, ["PCNR_AGE"])
    total_n = 100
    df = _make_rows(
        {"family": "F", "variable": "PCNR_AGE", "var_class": "CONT_NUM", "stat": "N",
         "level": "", "n": str(total_n), "value": "", "pct_total": "", "pct_nonmiss": "",
         "sort_order": "0", "small_cell_flag": "N", "date_kind": "", "display_format": ""},
        {"family": "F", "variable": "PCNR_AGE", "var_class": "CONT_NUM", "stat": "LEVEL",
         "level": "90", "n": "5", "value": "", "pct_total": "5.0", "pct_nonmiss": "5.0",
         "sort_order": "1", "small_cell_flag": "Y", "date_kind": "", "display_format": ""},
        {"family": "F", "variable": "PCNR_AGE", "var_class": "CONT_NUM", "stat": "LEVEL",
         "level": "91", "n": "5", "value": "", "pct_total": "5.0", "pct_nonmiss": "5.0",
         "sort_order": "2", "small_cell_flag": "Y", "date_kind": "", "display_format": ""},
        {"family": "F", "variable": "PCNR_AGE", "var_class": "CONT_NUM", "stat": "LEVEL",
         "level": "92", "n": "5", "value": "", "pct_total": "5.0", "pct_nonmiss": "5.0",
         "sort_order": "3", "small_cell_flag": "Y", "date_kind": "", "display_format": ""},
    )
    df2 = apply_age_topcode(df, docs_dir / "age_variables.csv")
    level_rows = df2[df2["stat"] == "LEVEL"]
    assert list(level_rows["level"].unique()) == ["90+"], "All >89 LEVEL rows must merge to '90+'"
    merged_n = pd.to_numeric(level_rows.loc[level_rows["level"] == "90+", "n"]).iloc[0]
    assert merged_n == 15, f"Merged n should be 15, got {merged_n}"

    cfg = _cfg(small_cell_mode="suppress", age_topcode=False)
    df3 = apply_suppression(df2, cfg)
    row_90plus = df3[(df3["stat"] == "LEVEL") & (df3["level"] == "90+")]
    assert len(row_90plus) == 1
    assert row_90plus.iloc[0]["n"] != "<11", "'90+' row with n=15 must NOT be suppressed"


# ---------------------------------------------------------------------------
# Small-cell suppression modes
# ---------------------------------------------------------------------------

def test_display_rules_suppress_mode():
    """small_cell_flag='Y' row -> n='<11', pct_total='', pct_nonmiss='' in suppress mode."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "Yes", "n": "5", "pct_total": "5.0", "pct_nonmiss": "5.0",
         "small_cell_flag": "Y", "sort_order": "1"},
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "No", "n": "95", "pct_total": "95.0", "pct_nonmiss": "95.0",
         "small_cell_flag": "N", "sort_order": "2"},
    )
    cfg = _cfg(small_cell_mode="suppress", age_topcode=False)
    out = apply_suppression(df, cfg)
    yes_row = out[(out["stat"] == "LEVEL") & (out["level"] == "Yes")].iloc[0]
    assert yes_row["n"] == "<11"
    assert yes_row["pct_total"] == ""
    assert yes_row["pct_nonmiss"] == ""


def test_display_rules_raw_mode():
    """small_cell_flag='Y' row -> n unchanged in raw mode."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "Yes", "n": "5", "pct_total": "5.0", "pct_nonmiss": "5.0",
         "small_cell_flag": "Y", "sort_order": "1"},
    )
    cfg = _cfg(small_cell_mode="raw", age_topcode=False)
    out = apply_suppression(df, cfg)
    yes_row = out[(out["stat"] == "LEVEL") & (out["level"] == "Yes")].iloc[0]
    assert yes_row["n"] == "5"
    assert yes_row["pct_total"] == "5.0"


def test_display_rules_asterisk_mode():
    """small_cell_flag='Y' row -> level gets '*' appended in asterisk mode."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "Yes", "n": "5", "pct_total": "5.0", "pct_nonmiss": "5.0",
         "small_cell_flag": "Y", "sort_order": "1"},
    )
    cfg = _cfg(small_cell_mode="asterisk", age_topcode=False)
    out = apply_suppression(df, cfg)
    yes_row = out[out["stat"] == "LEVEL"].iloc[0]
    assert yes_row["level"].endswith("*"), f"Expected level ending with '*', got {yes_row['level']!r}"
    assert yes_row["n"] == "5"


# ---------------------------------------------------------------------------
# Complementary suppression
# ---------------------------------------------------------------------------

def test_complementary_suppression_fires():
    """Exactly 1 primary-suppressed row -> next-smallest row also suppressed as 'suppressed'."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "A", "n": "5", "pct_total": "5.0", "pct_nonmiss": "5.0",
         "small_cell_flag": "Y", "sort_order": "1"},
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "B", "n": "15", "pct_total": "15.0", "pct_nonmiss": "15.0",
         "small_cell_flag": "N", "sort_order": "2"},
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "C", "n": "80", "pct_total": "80.0", "pct_nonmiss": "80.0",
         "small_cell_flag": "N", "sort_order": "3"},
    )
    cfg = _cfg(small_cell_mode="suppress", age_topcode=False)
    out = apply_suppression(df, cfg)
    b_row = out[(out["level"] == "B")].iloc[0]
    assert b_row["n"] == "suppressed", "Next-smallest must be complementary-suppressed"
    c_row = out[(out["level"] == "C")].iloc[0]
    assert c_row["n"] == "80", "Largest row must remain unsuppressed"


def test_complementary_suppression_no_fire_two_suppressed():
    """2 primary-suppressed rows -> no complementary suppression."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "A", "n": "5", "pct_total": "5.0", "pct_nonmiss": "5.0",
         "small_cell_flag": "Y", "sort_order": "1"},
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "B", "n": "8", "pct_total": "8.0", "pct_nonmiss": "8.0",
         "small_cell_flag": "Y", "sort_order": "2"},
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "C", "n": "87", "pct_total": "87.0", "pct_nonmiss": "87.0",
         "small_cell_flag": "N", "sort_order": "3"},
    )
    cfg = _cfg(small_cell_mode="suppress", age_topcode=False)
    out = apply_suppression(df, cfg)
    c_row = out[(out["level"] == "C")].iloc[0]
    assert c_row["n"] == "87", "No complementary suppression when 2 primary-suppressed rows exist"


def test_complementary_suppression_ties():
    """Two candidates tied at min n -> both are complementary-suppressed."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "A", "n": "5", "pct_total": "5.0", "pct_nonmiss": "5.0",
         "small_cell_flag": "Y", "sort_order": "1"},
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "B", "n": "20", "pct_total": "20.0", "pct_nonmiss": "20.0",
         "small_cell_flag": "N", "sort_order": "2"},
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "C", "n": "20", "pct_total": "20.0", "pct_nonmiss": "20.0",
         "small_cell_flag": "N", "sort_order": "3"},
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "D", "n": "55", "pct_total": "55.0", "pct_nonmiss": "55.0",
         "small_cell_flag": "N", "sort_order": "4"},
    )
    cfg = _cfg(small_cell_mode="suppress", age_topcode=False)
    out = apply_suppression(df, cfg)
    b_row = out[(out["level"] == "B")].iloc[0]
    c_row = out[(out["level"] == "C")].iloc[0]
    assert b_row["n"] == "suppressed", "Both tied candidates must be complementary-suppressed"
    assert c_row["n"] == "suppressed", "Both tied candidates must be complementary-suppressed"


# ---------------------------------------------------------------------------
# NUM_TEXT unparsed level masking
# ---------------------------------------------------------------------------

def test_numtext_unparsed_level_masked():
    """NUM_TEXT LEVEL rows with n<=10 have level replaced with 'other text entry' before suppress."""
    df = _make_rows(
        {"variable": "VAR_T", "var_class": "NUM_TEXT", "stat": "LEVEL",
         "level": "some text value", "n": "5", "pct_total": "5.0", "pct_nonmiss": "5.0",
         "small_cell_flag": "Y", "sort_order": "1"},
        {"variable": "VAR_T", "var_class": "NUM_TEXT", "stat": "LEVEL",
         "level": "big text value", "n": "95", "pct_total": "95.0", "pct_nonmiss": "95.0",
         "small_cell_flag": "N", "sort_order": "2"},
    )
    cfg = _cfg(small_cell_mode="suppress", age_topcode=False)
    out = apply_suppression(df, cfg)
    small_row = out[out["sort_order"] == "1"].iloc[0]
    assert small_row["level"] == "other text entry", "Small NUM_TEXT LEVEL must be masked"
    big_row = out[out["sort_order"] == "2"].iloc[0]
    assert big_row["level"] == "big text value", "Large NUM_TEXT LEVEL must not be masked"


# ---------------------------------------------------------------------------
# Age topcode
# ---------------------------------------------------------------------------

def test_age_topcode_level_gt89(tmp_path):
    """LEVEL row level='91' for age var -> becomes '90+'."""
    docs_dir = _make_age_csv(tmp_path, ["PCNR_AGE"])
    df = _make_rows(
        {"variable": "PCNR_AGE", "var_class": "CONT_NUM", "stat": "LEVEL",
         "level": "91", "n": "20", "small_cell_flag": "N"},
    )
    out = apply_age_topcode(df, docs_dir / "age_variables.csv")
    assert out.iloc[0]["level"] == "90+"


def test_age_topcode_level_aggregation(tmp_path):
    """Three LEVEL rows n=3/4/5 (level='90'/'91'/'92') -> merged '90+' n=12; small_cell_flag='N'."""
    docs_dir = _make_age_csv(tmp_path, ["PCNR_AGE"])
    df = _make_rows(
        {"variable": "PCNR_AGE", "var_class": "CONT_NUM", "stat": "N",
         "n": "100", "sort_order": "0", "small_cell_flag": "N"},
        {"variable": "PCNR_AGE", "var_class": "CONT_NUM", "stat": "LEVEL",
         "level": "90", "n": "3", "pct_total": "3.0", "pct_nonmiss": "3.0",
         "sort_order": "1", "small_cell_flag": "Y"},
        {"variable": "PCNR_AGE", "var_class": "CONT_NUM", "stat": "LEVEL",
         "level": "91", "n": "4", "pct_total": "4.0", "pct_nonmiss": "4.0",
         "sort_order": "2", "small_cell_flag": "Y"},
        {"variable": "PCNR_AGE", "var_class": "CONT_NUM", "stat": "LEVEL",
         "level": "92", "n": "5", "pct_total": "5.0", "pct_nonmiss": "5.0",
         "sort_order": "3", "small_cell_flag": "Y"},
    )
    out = apply_age_topcode(df, docs_dir / "age_variables.csv")
    level_rows = out[out["stat"] == "LEVEL"]
    assert len(level_rows) == 1, "Three >89 LEVEL rows must merge into one"
    merged = level_rows.iloc[0]
    assert merged["level"] == "90+"
    assert int(float(merged["n"])) == 12
    assert merged["small_cell_flag"] == "N", "n=12>10 -> small_cell_flag='N'"


def test_age_topcode_stat_row_value(tmp_path):
    """MIN stat row with value='95.0' for age var -> value becomes '90+'."""
    docs_dir = _make_age_csv(tmp_path, ["PCNR_AGE"])
    df = _make_rows(
        {"variable": "PCNR_AGE", "var_class": "CONT_NUM", "stat": "MIN",
         "level": "", "n": "", "value": "95.0", "small_cell_flag": ""},
    )
    out = apply_age_topcode(df, docs_dir / "age_variables.csv")
    assert out.iloc[0]["value"] == "90+"


def test_age_topcode_mean_unchanged(tmp_path):
    """MEAN stat row with value='72.3' for age var -> value unchanged."""
    docs_dir = _make_age_csv(tmp_path, ["PCNR_AGE"])
    df = _make_rows(
        {"variable": "PCNR_AGE", "var_class": "CONT_NUM", "stat": "MEAN",
         "level": "", "n": "", "value": "72.3", "small_cell_flag": ""},
    )
    out = apply_age_topcode(df, docs_dir / "age_variables.csv")
    assert out.iloc[0]["value"] == "72.3"


def test_age_topcode_non_age_var(tmp_path):
    """LEVEL row level='91' for non-age var -> unchanged."""
    docs_dir = _make_age_csv(tmp_path, ["PCNR_AGE"])
    df = _make_rows(
        {"variable": "PCNR_OTHER", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "91", "n": "20", "small_cell_flag": "N"},
    )
    out = apply_age_topcode(df, docs_dir / "age_variables.csv")
    assert out.iloc[0]["level"] == "91"


# ---------------------------------------------------------------------------
# VARIABLES sheet derivation
# ---------------------------------------------------------------------------

def _run_build_variables(profile_df, namemap_rows=None, total_rows=None):
    """Helper: call build_variables_sheet with minimal name_map data."""
    if namemap_rows is None:
        namemap_rows = []
    namemap_df = pd.DataFrame(
        namemap_rows,
        columns=["s_code", "col_order", "sas_name", "original_header", "trunc_flag",
                 "status", "group_id", "pcnr_name", "sas_type", "length"],
    ) if not namemap_rows else pd.DataFrame(namemap_rows)
    cfg = _cfg(small_cell_mode="suppress", age_topcode=False)
    dist_row_map: dict[str, int] = {}
    return build_variables_sheet(profile_df, namemap_df, dist_row_map, cfg, total_rows=total_rows)


def test_variables_n_nonmiss_low_card():
    """LOW_CARD: n_nonmiss = total_rows - MISSING row n."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "Yes", "n": "93", "pct_total": "93.0", "pct_nonmiss": "100.0",
         "small_cell_flag": "N", "sort_order": "1"},
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "MISSING",
         "level": "", "n": "7", "pct_total": "7.0", "pct_nonmiss": "",
         "small_cell_flag": "N", "sort_order": "99"},
    )
    headers, rows = _run_build_variables(df, total_rows=100)
    row = rows[0]
    n_idx = headers.index("n_nonmiss")
    assert row[n_idx] == 93


def test_variables_n_nonmiss_cont_num():
    """CONT_NUM: n_nonmiss = N stat row n."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "N",
         "n": "95", "sort_order": "0", "small_cell_flag": "N"},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "MEAN",
         "value": "50.0", "sort_order": "1", "small_cell_flag": ""},
    )
    headers, rows = _run_build_variables(df, total_rows=100)
    row = rows[0]
    n_idx = headers.index("n_nonmiss")
    assert row[n_idx] == 95


def test_variables_n_nonmiss_num_text():
    """NUM_TEXT: n_nonmiss = N + N_UNPARSED."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "NUM_TEXT", "stat": "N",
         "n": "80", "sort_order": "0", "small_cell_flag": "N"},
        {"variable": "VAR1", "var_class": "NUM_TEXT", "stat": "N_UNPARSED",
         "n": "12", "sort_order": "1", "small_cell_flag": "N"},
    )
    headers, rows = _run_build_variables(df, total_rows=100)
    row = rows[0]
    n_idx = headers.index("n_nonmiss")
    assert row[n_idx] == 92


def test_variables_n_distinct_low_card():
    """LOW_CARD: n_distinct = count of LEVEL rows."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "A", "n": "30", "sort_order": "1", "small_cell_flag": "N"},
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "B", "n": "40", "sort_order": "2", "small_cell_flag": "N"},
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "C", "n": "30", "sort_order": "3", "small_cell_flag": "N"},
    )
    headers, rows = _run_build_variables(df, total_rows=100)
    row = rows[0]
    nd_idx = headers.index("n_distinct")
    assert row[nd_idx] == 3


def test_variables_n_distinct_high_card():
    """HIGH_CARD: n_distinct from N_DISTINCT stat row."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "HIGH_CARD", "stat": "N_DISTINCT",
         "n": "350", "sort_order": "0", "small_cell_flag": "N"},
        {"variable": "VAR1", "var_class": "HIGH_CARD", "stat": "N",
         "n": "1000", "sort_order": "1", "small_cell_flag": "N"},
    )
    headers, rows = _run_build_variables(df, total_rows=1000)
    row = rows[0]
    nd_idx = headers.index("n_distinct")
    assert row[nd_idx] == 350


def test_variables_n_distinct_cont_num():
    """CONT_NUM: n_distinct is blank."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "N",
         "n": "95", "sort_order": "0", "small_cell_flag": "N"},
    )
    headers, rows = _run_build_variables(df, total_rows=100)
    row = rows[0]
    nd_idx = headers.index("n_distinct")
    assert row[nd_idx] == "" or row[nd_idx] is None


def test_variables_all_missing_flag():
    """n_nonmiss==0 -> all_missing='Y'."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "MISSING",
         "level": "", "n": "100", "pct_total": "100.0", "pct_nonmiss": "",
         "small_cell_flag": "N", "sort_order": "99"},
    )
    headers, rows = _run_build_variables(df, total_rows=100)
    row = rows[0]
    am_idx = headers.index("all_missing")
    assert row[am_idx] == "Y"


def test_variables_suppressed_missing_row_leak():
    """
    MISSING row suppressed (n='<11') -> VARIABLES n_nonmiss shows '<11 missing', not an integer.
    """
    df = _make_rows(
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "Yes", "n": "95", "pct_total": "95.0", "pct_nonmiss": "100.0",
         "small_cell_flag": "N", "sort_order": "1"},
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "MISSING",
         "level": "", "n": "5", "pct_total": "5.0", "pct_nonmiss": "",
         "small_cell_flag": "Y", "sort_order": "99"},
    )
    cfg = _cfg(small_cell_mode="suppress", age_topcode=False)
    suppressed_df = apply_suppression(df, cfg)
    dist_row_map: dict[str, int] = {}
    namemap_df = pd.DataFrame(columns=["s_code", "col_order", "sas_name", "original_header",
                                        "trunc_flag", "status", "group_id", "pcnr_name",
                                        "sas_type", "length"])
    headers, rows = build_variables_sheet(suppressed_df, namemap_df, dist_row_map, cfg, total_rows=100)
    row = rows[0]
    n_idx = headers.index("n_nonmiss")
    assert row[n_idx] == "<11 missing", f"Expected '<11 missing', got {row[n_idx]!r}"


# ---------------------------------------------------------------------------
# Stat-class suppression
# ---------------------------------------------------------------------------

def test_stat_group_suppress_n_missing():
    """
    CONT_NUM var, N=100, N_MISSING=1 (primary-suppressed).
    N_MISSING displayed as '<11'; N complementary -> 'suppressed'; pct columns blank.
    MEAN/SD etc. UNCHANGED. VARIABLES summary normal.
    """
    df = _make_rows(
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "N",
         "n": "100", "pct_total": "", "pct_nonmiss": "", "sort_order": "0", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "N_MISSING",
         "n": "1", "pct_total": "1.0", "pct_nonmiss": "", "sort_order": "0", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "MEAN",
         "n": "", "value": "50.5", "sort_order": "1", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "SD",
         "n": "", "value": "10.2", "sort_order": "2", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "MEDIAN",
         "n": "", "value": "51.0", "sort_order": "3", "small_cell_flag": ""},
    )
    cfg = _cfg(small_cell_mode="suppress", age_topcode=False)
    out = apply_suppression(df, cfg)

    n_miss_row = out[out["stat"] == "N_MISSING"].iloc[0]
    assert n_miss_row["n"] == "<11", "N_MISSING primary-suppressed -> '<11'"
    assert n_miss_row["pct_total"] == ""

    n_row = out[out["stat"] == "N"].iloc[0]
    assert n_row["n"] == "suppressed", "N complementary-suppressed -> 'suppressed'"
    assert n_row["pct_total"] == ""

    mean_row = out[out["stat"] == "MEAN"].iloc[0]
    assert mean_row["value"] == "50.5", "MEAN unchanged when N is complementary-suppressed"

    median_row = out[out["stat"] == "MEDIAN"].iloc[0]
    assert median_row["value"] == "51.0", "MEDIAN unchanged when N is complementary-suppressed"


def test_stat_group_suppress_n_unparsed():
    """NUM_TEXT var, N_UNPARSED=5 -> '<11'; N='suppressed' (complementary). Stats unchanged."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "NUM_TEXT", "stat": "N",
         "n": "75", "pct_total": "", "pct_nonmiss": "", "sort_order": "0", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "NUM_TEXT", "stat": "N_UNPARSED",
         "n": "5", "pct_total": "5.0", "pct_nonmiss": "", "sort_order": "1", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "NUM_TEXT", "stat": "MEAN",
         "n": "", "value": "30.0", "sort_order": "2", "small_cell_flag": ""},
    )
    cfg = _cfg(small_cell_mode="suppress", age_topcode=False)
    out = apply_suppression(df, cfg)

    unp_row = out[out["stat"] == "N_UNPARSED"].iloc[0]
    assert unp_row["n"] == "<11"
    assert unp_row["pct_total"] == ""

    n_row = out[out["stat"] == "N"].iloc[0]
    assert n_row["n"] == "suppressed"

    mean_row = out[out["stat"] == "MEAN"].iloc[0]
    assert mean_row["value"] == "30.0", "MEAN unchanged"


def test_stat_group_suppress_n_itself():
    """
    CONT_NUM var, N=3 (primary-suppressed).
    N -> '<11'; all stat value cells blank; VARIABLES summary 'fewer than 11 values'; n_nonmiss='<11'.
    """
    df = _make_rows(
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "N",
         "n": "3", "pct_total": "", "pct_nonmiss": "", "sort_order": "0", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "MEAN",
         "n": "", "value": "50.5", "sort_order": "1", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "SD",
         "n": "", "value": "10.2", "sort_order": "2", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "MEDIAN",
         "n": "", "value": "51.0", "sort_order": "3", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "Q1",
         "n": "", "value": "45.0", "sort_order": "4", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "Q3",
         "n": "", "value": "55.0", "sort_order": "5", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "MIN",
         "n": "", "value": "10.0", "sort_order": "6", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "MAX",
         "n": "", "value": "90.0", "sort_order": "7", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "RANGE_MIN",
         "n": "", "value": "10.0", "sort_order": "8", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "RANGE_MAX",
         "n": "", "value": "90.0", "sort_order": "9", "small_cell_flag": ""},
    )
    cfg = _cfg(small_cell_mode="suppress", age_topcode=False)
    out = apply_suppression(df, cfg)

    n_row = out[out["stat"] == "N"].iloc[0]
    assert n_row["n"] == "<11", "N=3 primary-suppressed -> '<11'"

    for stat_name in ["MEAN", "SD", "MEDIAN", "Q1", "Q3", "MIN", "MAX", "RANGE_MIN", "RANGE_MAX"]:
        stat_row = out[out["stat"] == stat_name].iloc[0]
        assert stat_row["value"] == "", f"{stat_name} must be blanked when N is primary-suppressed"

    # VARIABLES sheet must reflect suppressed N
    dist_row_map: dict[str, int] = {}
    namemap_df = pd.DataFrame(columns=["s_code", "col_order", "sas_name", "original_header",
                                        "trunc_flag", "status", "group_id", "pcnr_name",
                                        "sas_type", "length"])
    headers, rows = build_variables_sheet(out, namemap_df, dist_row_map, cfg, total_rows=100)
    var_row = rows[0]
    n_idx = headers.index("n_nonmiss")
    summary_idx = headers.index("summary")
    pct_idx = headers.index("pct_missing")
    assert var_row[n_idx] == "<11", f"n_nonmiss should be '<11', got {var_row[n_idx]!r}"
    assert var_row[summary_idx] == "fewer than 11 values"
    assert var_row[pct_idx] == ""


def test_stat_group_complementary_suppress():
    """N_MISSING=4 (primary); exactly one primary -> N='suppressed'. VARIABLES summary unchanged."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "N",
         "n": "100", "pct_total": "", "pct_nonmiss": "", "sort_order": "0", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "N_MISSING",
         "n": "4", "pct_total": "4.0", "pct_nonmiss": "", "sort_order": "1", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "MEAN",
         "n": "", "value": "50.0", "sort_order": "2", "small_cell_flag": ""},
    )
    cfg = _cfg(small_cell_mode="suppress", age_topcode=False)
    out = apply_suppression(df, cfg)

    n_row = out[out["stat"] == "N"].iloc[0]
    assert n_row["n"] == "suppressed"

    dist_row_map: dict[str, int] = {}
    namemap_df = pd.DataFrame(columns=["s_code", "col_order", "sas_name", "original_header",
                                        "trunc_flag", "status", "group_id", "pcnr_name",
                                        "sas_type", "length"])
    headers, rows = build_variables_sheet(out, namemap_df, dist_row_map, cfg, total_rows=100)
    var_row = rows[0]
    summary_idx = headers.index("summary")
    assert var_row[summary_idx] != "fewer than 11 values", (
        "Summary should not be 'fewer than 11 values' when N is complementary-suppressed"
    )


# ---------------------------------------------------------------------------
# Summary text
# ---------------------------------------------------------------------------

def test_summary_text_low_card():
    """LOW_CARD fixture with 3 LEVEL rows -> summary matches 'Val X.X%' format."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "Yes", "n": "45", "pct_total": "45.2", "pct_nonmiss": "47.0",
         "small_cell_flag": "N", "sort_order": "1"},
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "No", "n": "52", "pct_total": "52.1", "pct_nonmiss": "53.0",
         "small_cell_flag": "N", "sort_order": "2"},
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "LEVEL",
         "level": "Unknown", "n": "3", "pct_total": "2.7", "pct_nonmiss": "0.0",
         "small_cell_flag": "N", "sort_order": "3"},
    )
    cfg = _cfg(small_cell_mode="suppress", age_topcode=False)
    dist_row_map: dict[str, int] = {}
    namemap_df = pd.DataFrame(columns=["s_code", "col_order", "sas_name", "original_header",
                                        "trunc_flag", "status", "group_id", "pcnr_name",
                                        "sas_type", "length"])
    headers, rows = build_variables_sheet(df, namemap_df, dist_row_map, cfg, total_rows=100)
    var_row = rows[0]
    summary_idx = headers.index("summary")
    summary = var_row[summary_idx]
    assert "%" in summary, f"LOW_CARD summary must contain '%', got: {summary!r}"
    assert "Yes" in summary or "No" in summary, f"Summary must include level names, got: {summary!r}"


def test_summary_text_all_missing():
    """LOW_CARD var with n_nonmiss==0 -> summary == 'No values (100% missing)'."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "LOW_CARD", "stat": "MISSING",
         "level": "", "n": "100", "pct_total": "100.0", "pct_nonmiss": "",
         "small_cell_flag": "N", "sort_order": "99"},
    )
    cfg = _cfg(small_cell_mode="suppress", age_topcode=False)
    dist_row_map: dict[str, int] = {}
    namemap_df = pd.DataFrame(columns=["s_code", "col_order", "sas_name", "original_header",
                                        "trunc_flag", "status", "group_id", "pcnr_name",
                                        "sas_type", "length"])
    headers, rows = build_variables_sheet(df, namemap_df, dist_row_map, cfg, total_rows=100)
    var_row = rows[0]
    summary_idx = headers.index("summary")
    assert var_row[summary_idx] == "No values (100% missing)"


def test_summary_text_cont_num():
    """CONT_NUM fixture -> summary starts with 'mean'."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "N",
         "n": "95", "sort_order": "0", "small_cell_flag": "N"},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "MEAN",
         "value": "50.5", "sort_order": "1", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "SD",
         "value": "10.2", "sort_order": "2", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "MEDIAN",
         "value": "51.0", "sort_order": "3", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "Q1",
         "value": "45.0", "sort_order": "4", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "Q3",
         "value": "55.0", "sort_order": "5", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "MIN",
         "value": "10.0", "sort_order": "6", "small_cell_flag": ""},
        {"variable": "VAR1", "var_class": "CONT_NUM", "stat": "MAX",
         "value": "90.0", "sort_order": "7", "small_cell_flag": ""},
    )
    cfg = _cfg(small_cell_mode="suppress", age_topcode=False)
    dist_row_map: dict[str, int] = {}
    namemap_df = pd.DataFrame(columns=["s_code", "col_order", "sas_name", "original_header",
                                        "trunc_flag", "status", "group_id", "pcnr_name",
                                        "sas_type", "length"])
    headers, rows = build_variables_sheet(df, namemap_df, dist_row_map, cfg, total_rows=100)
    var_row = rows[0]
    summary_idx = headers.index("summary")
    assert var_row[summary_idx].lower().startswith("mean")


def test_summary_text_suppressed_range():
    """SUPPRESSED var with display_format != 'LENGTH' -> summary contains 'range'."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "SUPPRESSED", "stat": "N",
         "n": "100", "sort_order": "0", "small_cell_flag": "", "display_format": ""},
        {"variable": "VAR1", "var_class": "SUPPRESSED", "stat": "RANGE_MIN",
         "value": "0.0", "sort_order": "1", "small_cell_flag": "", "display_format": ""},
        {"variable": "VAR1", "var_class": "SUPPRESSED", "stat": "RANGE_MAX",
         "value": "100.0", "sort_order": "2", "small_cell_flag": "", "display_format": ""},
    )
    cfg = _cfg(small_cell_mode="suppress", age_topcode=False)
    dist_row_map: dict[str, int] = {}
    namemap_df = pd.DataFrame(columns=["s_code", "col_order", "sas_name", "original_header",
                                        "trunc_flag", "status", "group_id", "pcnr_name",
                                        "sas_type", "length"])
    headers, rows = build_variables_sheet(df, namemap_df, dist_row_map, cfg, total_rows=100)
    var_row = rows[0]
    summary_idx = headers.index("summary")
    assert "range" in var_row[summary_idx].lower()


def test_summary_text_suppressed_length():
    """SUPPRESSED var with display_format=='LENGTH' -> summary contains 'length'."""
    df = _make_rows(
        {"variable": "VAR1", "var_class": "SUPPRESSED", "stat": "N",
         "n": "100", "sort_order": "0", "small_cell_flag": "", "display_format": "LENGTH"},
        {"variable": "VAR1", "var_class": "SUPPRESSED", "stat": "RANGE_MIN",
         "value": "1.0", "sort_order": "1", "small_cell_flag": "", "display_format": "LENGTH"},
        {"variable": "VAR1", "var_class": "SUPPRESSED", "stat": "RANGE_MAX",
         "value": "50.0", "sort_order": "2", "small_cell_flag": "", "display_format": "LENGTH"},
    )
    cfg = _cfg(small_cell_mode="suppress", age_topcode=False)
    dist_row_map: dict[str, int] = {}
    namemap_df = pd.DataFrame(columns=["s_code", "col_order", "sas_name", "original_header",
                                        "trunc_flag", "status", "group_id", "pcnr_name",
                                        "sas_type", "length"])
    headers, rows = build_variables_sheet(df, namemap_df, dist_row_map, cfg, total_rows=100)
    var_row = rows[0]
    summary_idx = headers.index("summary")
    assert "length" in var_row[summary_idx].lower()


# ---------------------------------------------------------------------------
# Workbook structure
# ---------------------------------------------------------------------------

def _minimal_write_workbook(tmp_path):
    """Write a workbook from minimal fixture data; return the output path."""
    exchange_dir = tmp_path / "exchange"
    exchange_dir.mkdir()
    docs_dir = tmp_path / "docs"
    docs_dir.mkdir()
    out_path = tmp_path / "out" / "test.xlsx"

    # age_variables.csv — at least one row so hard-fail does not trigger
    (docs_dir / "age_variables.csv").write_text("variable,reason\nPCNR_PLACEHOLDER,test\n")

    # profile_PCNR.csv — two variables, one LOW_CARD, one CONT_NUM
    profile_rows = [
        "family,variable,var_class,stat,level,n,value,pct_total,pct_nonmiss,"
        "sort_order,small_cell_flag,date_kind,display_format",
        "F,PCNR_A,LOW_CARD,LEVEL,Yes,80,,80.0,88.9,1,N,,",
        "F,PCNR_A,LOW_CARD,LEVEL,No,10,,10.0,11.1,2,N,,",
        "F,PCNR_A,LOW_CARD,MISSING,,10,,10.0,,99,N,,",
        "F,PCNR_B,CONT_NUM,N,,95,,,, 0,N,,",
        "F,PCNR_B,CONT_NUM,MEAN,,,50.5,,,1,,,",
        "F,PCNR_B,CONT_NUM,SD,,,10.2,,,2,,,",
        "F,PCNR_B,CONT_NUM,MEDIAN,,,51.0,,,3,,,",
        "F,PCNR_B,CONT_NUM,Q1,,,45.0,,,4,,,",
        "F,PCNR_B,CONT_NUM,Q3,,,55.0,,,5,,,",
        "F,PCNR_B,CONT_NUM,MIN,,,10.0,,,6,,,",
        "F,PCNR_B,CONT_NUM,MAX,,,90.0,,,7,,,",
    ]
    (exchange_dir / "profile_PCNR.csv").write_text("\n".join(profile_rows))

    # name_map_dict.csv
    nm_header = "s_code,col_order,sas_name,original_header,trunc_flag,status,group_id,pcnr_name,sas_type,length"
    nm_row = "s01,1,PCNR_A,A,0,PCNR,1,PCNR_A,C,32"
    nm_row2 = "s01,2,PCNR_B,B,0,PCNR,2,PCNR_B,N,8"
    (exchange_dir / "name_map_dict.csv").write_text(f"{nm_header}\n{nm_row}\n{nm_row2}")

    # source_inventory.csv
    si_header = "s_code,filename,row_count,col_count,sha256,mod_date"
    si_row = "s01,file1.csv,100,5,abc123def456abc123def456abc123def456abc123def456abc123def456abc1,2026-01-01"
    (exchange_dir / "source_inventory.csv").write_text(f"{si_header}\n{si_row}")

    cfg = _cfg(small_cell_mode="suppress", age_topcode=True,
               age_variables_file="age_variables.csv")
    counts = write_workbook(exchange_dir, out_path, docs_dir, config=cfg)
    return out_path, counts


def test_sheet_order(tmp_path):
    """write_workbook from minimal fixture -> wb.sheetnames == SHEET_ORDER."""
    from openpyxl import load_workbook as lw
    out_path, _ = _minimal_write_workbook(tmp_path)
    wb = lw(out_path)
    assert wb.sheetnames == SHEET_ORDER


def test_hyperlink_is_internal(tmp_path):
    """pcnr_name cell in VARIABLES has hyperlink with location starting 'DISTRIBUTIONS!'."""
    from openpyxl import load_workbook as lw
    out_path, _ = _minimal_write_workbook(tmp_path)
    wb = lw(out_path)
    ws = wb["VARIABLES"]
    # Find pcnr_name column
    header = [c.value for c in ws[1]]
    col_idx = header.index("pcnr_name") + 1
    found_internal = False
    for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
        cell = row[0]
        if cell.hyperlink and cell.hyperlink.location:
            assert cell.hyperlink.location.startswith("DISTRIBUTIONS!")
            found_internal = True
            break
    assert found_internal, "At least one pcnr_name cell must have an internal hyperlink"


def test_name_map_sheet_columns(tmp_path):
    """NAME_MAP sheet header row == expected 10 columns."""
    from openpyxl import load_workbook as lw
    out_path, _ = _minimal_write_workbook(tmp_path)
    wb = lw(out_path)
    ws = wb["NAME_MAP"]
    header = [c.value for c in ws[1]]
    assert header == [
        "s_code", "col_order", "sas_name", "original_header", "trunc_flag",
        "status", "group_id", "pcnr_name", "sas_type", "length",
    ]


def test_sources_no_notes(tmp_path):
    """build_sources_sheet works when source_notes.csv is absent; description column present."""
    inventory_df = pd.DataFrame([{
        "s_code": "s01", "filename": "file1.csv", "row_count": "100",
        "col_count": "5", "sha256": "abc123def456xyz789", "mod_date": "2026-01-01",
    }])
    docs_dir = tmp_path / "docs"
    docs_dir.mkdir()
    # Do NOT create source_notes.csv
    headers, rows = build_sources_sheet(inventory_df, docs_dir)
    assert "description" in headers
    # Should not raise


def test_distributions_row_count(tmp_path):
    """DISTRIBUTIONS sheet data row count == len(final_display_df) after transforms."""
    from openpyxl import load_workbook as lw
    out_path, counts = _minimal_write_workbook(tmp_path)
    wb = lw(out_path)
    ws = wb["DISTRIBUTIONS"]
    data_rows = ws.max_row - 1
    assert data_rows == counts["DISTRIBUTIONS"], (
        f"Expected {counts['DISTRIBUTIONS']} data rows, got {data_rows}"
    )


def test_write_workbook_returns_dict(tmp_path):
    """write_workbook returns dict with keys VARIABLES, DISTRIBUTIONS, NAME_MAP."""
    _, counts = _minimal_write_workbook(tmp_path)
    assert "VARIABLES" in counts
    assert "DISTRIBUTIONS" in counts
    assert "NAME_MAP" in counts


def test_placeholder_sheets(tmp_path):
    """Each placeholder sheet has A1 == correct pending string."""
    from openpyxl import load_workbook as lw
    out_path, _ = _minimal_write_workbook(tmp_path)
    wb = lw(out_path)
    expected = {
        "PCNC_": "(source pending — Phase 8)",
        "PCND_": "(source pending — Phase 13)",
        "CLK_": "(source pending — Phase 9)",
        "LAB_": "(source pending — Phase 10)",
        "MED_": "(source pending — Phase 11)",
        "INTRA_": "(source pending — Phase 12)",
    }
    for sheet_name, expected_text in expected.items():
        ws = wb[sheet_name]
        a1 = ws["A1"].value
        assert a1 == expected_text, f"{sheet_name}!A1: expected {expected_text!r}, got {a1!r}"


# ---------------------------------------------------------------------------
# No-leak scan
# ---------------------------------------------------------------------------

def test_no_suppressed_count_leaks(tmp_path):
    """
    Write workbook with suppressed cells (LEVEL-class and stat-class).
    Load all sheets. Assert no cell has a raw integer 1-10 matching a suppressed n.
    Also assert no suppressed level text (from NUM_TEXT masking) appears.
    Includes stat-class row (N_MISSING=3) to verify leak prevention.
    """
    from openpyxl import load_workbook as lw

    exchange_dir = tmp_path / "exchange"
    exchange_dir.mkdir()
    docs_dir = tmp_path / "docs"
    docs_dir.mkdir()
    out_path = tmp_path / "out2" / "test2.xlsx"

    (docs_dir / "age_variables.csv").write_text("variable,reason\nPCNR_PLACEHOLDER,test\n")

    # profile_PCNR.csv with:
    # - LOW_CARD var: LEVEL row n=7 (small-cell), LEVEL row n=93
    # - NUM_TEXT var: LEVEL row n=5 (small-cell + text masking), N=80, N_MISSING=3 (stat-class primary)
    profile_rows = (
        "family,variable,var_class,stat,level,n,value,pct_total,pct_nonmiss,"
        "sort_order,small_cell_flag,date_kind,display_format\n"
        "F,PCNR_LC,LOW_CARD,LEVEL,SuppressedVal,7,,7.0,7.5,1,Y,,\n"
        "F,PCNR_LC,LOW_CARD,LEVEL,BigVal,93,,93.0,93.0,2,N,,\n"
        "F,PCNR_NT,NUM_TEXT,N,,80,,,, 0,N,,\n"
        "F,PCNR_NT,NUM_TEXT,N_MISSING,,3,,,, 0,N,,\n"
        "F,PCNR_NT,NUM_TEXT,LEVEL,SecretText,5,,5.0,6.0,1,Y,,\n"
        "F,PCNR_NT,NUM_TEXT,LEVEL,BigText,75,,75.0,90.0,2,N,,\n"
    )
    (exchange_dir / "profile_PCNR.csv").write_text(profile_rows)

    nm_header = "s_code,col_order,sas_name,original_header,trunc_flag,status,group_id,pcnr_name,sas_type,length"
    (exchange_dir / "name_map_dict.csv").write_text(
        f"{nm_header}\ns01,1,PCNR_LC,LC,0,PCNR,1,PCNR_LC,C,32\ns01,2,PCNR_NT,NT,0,PCNR,2,PCNR_NT,C,32"
    )
    si_header = "s_code,filename,row_count,col_count,sha256,mod_date"
    si_row = "s01,file1.csv,100,5,abc123def456abc123def456abc123def456abc123def456abc123def456abc1,2026-01-01"
    (exchange_dir / "source_inventory.csv").write_text(f"{si_header}\n{si_row}")

    cfg = _cfg(small_cell_mode="suppress", age_topcode=True,
               age_variables_file="age_variables.csv")
    write_workbook(exchange_dir, out_path, docs_dir, config=cfg)

    wb = lw(out_path)
    forbidden_ints = {7, 5, 3}  # suppressed n values from fixture
    forbidden_text = "SecretText"

    # Only scan sheets that contain profile data; SOURCES/NAME_MAP/KEY/QC may legitimately
    # contain integers like col_count=5 that are unrelated to suppressed profile counts.
    profile_sheets = {"VARIABLES", "DISTRIBUTIONS"}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if v is None:
                    continue
                # Check no masked level text leaks (all sheets)
                if isinstance(v, str) and forbidden_text in v:
                    pytest.fail(
                        f"Suppressed level text '{forbidden_text}' leaked in "
                        f"{sheet_name}!{cell.coordinate}"
                    )
                if sheet_name not in profile_sheets:
                    continue
                # Check no raw integer leaks in profile sheets
                if isinstance(v, int) and v in forbidden_ints:
                    pytest.fail(
                        f"Suppressed count {v} leaked in {sheet_name}!{cell.coordinate}"
                    )
                if isinstance(v, str):
                    try:
                        as_int = int(float(v))
                        if as_int in forbidden_ints and str(as_int) == v.strip():
                            pytest.fail(
                                f"Suppressed count string '{v}' leaked in "
                                f"{sheet_name}!{cell.coordinate}"
                            )
                    except (ValueError, TypeError):
                        pass


# ---------------------------------------------------------------------------
# pandas 3.x compat: no FutureWarning / SettingWithCopyWarning


def test_no_pandas_warnings_from_write_workbook(tmp_path):
    """
    write_workbook must emit no FutureWarning or SettingWithCopyWarning.
    Covers pandas 3.x copy-on-write and new string dtype behaviour.
    """
    import warnings

    exchange_dir = tmp_path / "exchange"
    exchange_dir.mkdir()
    docs_dir = tmp_path / "docs"
    docs_dir.mkdir()
    out_path = tmp_path / "out3" / "test3.xlsx"

    (docs_dir / "age_variables.csv").write_text("variable,reason\nPCNR_Age_at_Encounter,age\n")

    profile_rows = (
        "family,variable,var_class,stat,level,n,value,pct_total,pct_nonmiss,"
        "sort_order,small_cell_flag,date_kind,display_format\n"
        # LOW_CARD: one small-cell row
        "F,PCNR_LC,LOW_CARD,LEVEL,SmallVal,5,,5.0,5.5,1,Y,,\n"
        "F,PCNR_LC,LOW_CARD,LEVEL,BigVal,95,,95.0,95.0,2,N,,\n"
        # CONT_NUM: age variable with LEVEL row >89 and small N_MISSING
        "F,PCNR_Age_at_Encounter,CONT_NUM,N,,100,,,, 0,,,\n"
        "F,PCNR_Age_at_Encounter,CONT_NUM,N_MISSING,,3,,,, 0,,,\n"
        "F,PCNR_Age_at_Encounter,CONT_NUM,MEAN,,,72.3,,,0,,,\n"
        "F,PCNR_Age_at_Encounter,CONT_NUM,MIN,,,91.0,,,0,,,\n"
        # NUM_TEXT: text-mask trigger
        "F,PCNR_NT,NUM_TEXT,N,,80,,,, 0,,,\n"
        "F,PCNR_NT,NUM_TEXT,LEVEL,SecretText,5,,5.0,6.0,1,Y,,\n"
    )
    (exchange_dir / "profile_PCNR.csv").write_text(profile_rows)

    nm_header = "s_code,col_order,sas_name,original_header,trunc_flag,status,group_id,pcnr_name,sas_type,length"
    (exchange_dir / "name_map_dict.csv").write_text(
        f"{nm_header}\n"
        "s01,1,PCNR_LC,LC,0,PCNR,1,PCNR_LC,C,32\n"
        "s01,2,PCNR_Age_at_Encounter,Age,0,PCNR,2,PCNR_Age_at_Encounter,N,8\n"
        "s01,3,PCNR_NT,NT,0,PCNR,3,PCNR_NT,C,32\n"
    )
    si_header = "s_code,filename,row_count,col_count,sha256,mod_date"
    (exchange_dir / "source_inventory.csv").write_text(
        f"{si_header}\ns01,file1.csv,100,5,{'a' * 64},2026-01-01\n"
    )

    cfg = _cfg(small_cell_mode="suppress", age_topcode=True,
               age_variables_file="age_variables.csv")

    with warnings.catch_warnings():
        warnings.simplefilter("error", FutureWarning)
        try:
            from pandas.errors import SettingWithCopyWarning
            warnings.simplefilter("error", SettingWithCopyWarning)
        except ImportError:
            pass  # pandas version without this warning class
        # Must not raise
        write_workbook(exchange_dir, out_path, docs_dir, config=cfg)


# ---------------------------------------------------------------------------
# Bug 1: QC n_all_missing_vars
# ---------------------------------------------------------------------------

def test_qc_n_all_missing_vars(tmp_path):
    """build_qc_sheet receives n_all_missing from build_variables_sheet output.

    Two CONT_NUM variables both have N=0 (all-missing). QC must report
    n_all_missing_vars=2, not 0.
    """
    import openpyxl

    profile_rows = (
        "family,variable,var_class,stat,level,n,value,pct_total,pct_nonmiss,"
        "sort_order,small_cell_flag,date_kind,display_format\n"
        "F,PCNR_A,CONT_NUM,N,,0,,,, 0,,,\n"
        "F,PCNR_A,CONT_NUM,N_MISSING,,100,,,, 0,,,\n"
        "F,PCNR_B,CONT_NUM,N,,0,,,, 0,,,\n"
        "F,PCNR_B,CONT_NUM,N_MISSING,,100,,,, 0,,,\n"
    )
    profile_df = pd.read_csv(
        __import__("io").StringIO(profile_rows), dtype=str, keep_default_na=False
    )
    namemap_df = pd.DataFrame(columns=["s_code", "pcnr_name", "original_header", "status"])
    cfg = _cfg()

    _, var_rows = build_variables_sheet(
        profile_df, namemap_df, {}, cfg, total_rows=100
    )
    n_all_missing = sum(1 for r in var_rows if r[8] == "Y")
    assert n_all_missing == 2, f"Expected 2 all-missing vars, got {n_all_missing}"

    wb = openpyxl.Workbook()
    ws = wb.active
    build_qc_sheet(profile_df, ws, n_all_missing)

    # Find the n_all_missing_vars row
    result = None
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == "n_all_missing_vars":
            result = row[1]
            break
    assert result == 2, f"QC n_all_missing_vars expected 2, got {result!r}"


# ---------------------------------------------------------------------------
# Bug 2: NUM_TEXT N_UNPARSED primary-suppressed, N visible
# ---------------------------------------------------------------------------

def test_num_text_n_unparsed_primary_suppressed_n_visible():
    """When N=1460 (visible) and N_UNPARSED=<11 (primary-suppressed):
    - n_nonmiss must equal 1460 (not '<11')
    - summary must NOT say 'fewer than 11 values'
    - summary must include '<11 non-numeric entries'
    - pct_missing must be blank (N_MISSING is complementary-suppressed)
    """
    profile_rows = _make_rows(
        {"variable": "PCNR_X", "var_class": "NUM_TEXT", "stat": "N", "n": "1460", "small_cell_flag": "N"},
        {"variable": "PCNR_X", "var_class": "NUM_TEXT", "stat": "N_UNPARSED", "n": "<11", "small_cell_flag": "Y"},
        {"variable": "PCNR_X", "var_class": "NUM_TEXT", "stat": "N_MISSING", "n": "suppressed", "small_cell_flag": "N"},
        {"variable": "PCNR_X", "var_class": "NUM_TEXT", "stat": "MEAN", "value": "5.3", "small_cell_flag": "N"},
        {"variable": "PCNR_X", "var_class": "NUM_TEXT", "stat": "SD", "value": "1.2", "small_cell_flag": "N"},
        {"variable": "PCNR_X", "var_class": "NUM_TEXT", "stat": "MEDIAN", "value": "5.0", "small_cell_flag": "N"},
    )
    namemap_df = pd.DataFrame(columns=["s_code", "pcnr_name", "original_header", "status"])
    cfg = _cfg()

    _, var_rows = build_variables_sheet(
        profile_rows, namemap_df, {}, cfg, total_rows=1466
    )
    assert var_rows, "Expected at least one VARIABLES row"
    row = var_rows[0]

    n_nonmiss = row[5]   # index 5 = n_nonmiss
    pct_missing = row[6]  # index 6 = pct_missing
    summary = row[9]      # index 9 = summary

    assert n_nonmiss == 1460, f"n_nonmiss should be 1460, got {n_nonmiss!r}"
    assert pct_missing == "", f"pct_missing should be blank when N_MISSING suppressed, got {pct_missing!r}"
    assert "fewer than 11 values" not in str(summary), (
        f"summary must not say 'fewer than 11 values' when N is visible; got {summary!r}"
    )
    assert "<11 non-numeric entries" in str(summary), (
        f"summary must include '<11 non-numeric entries'; got {summary!r}"
    )


def test_num_text_n_primary_suppressed_still_fewer_than_11():
    """When N itself is primary-suppressed (<11), summary = 'fewer than 11 values'."""
    profile_rows = _make_rows(
        {"variable": "PCNR_Y", "var_class": "NUM_TEXT", "stat": "N", "n": "<11", "small_cell_flag": "Y"},
        {"variable": "PCNR_Y", "var_class": "NUM_TEXT", "stat": "N_UNPARSED", "n": "suppressed", "small_cell_flag": "N"},
        {"variable": "PCNR_Y", "var_class": "NUM_TEXT", "stat": "N_MISSING", "n": "1000", "small_cell_flag": "N"},
    )
    namemap_df = pd.DataFrame(columns=["s_code", "pcnr_name", "original_header", "status"])
    cfg = _cfg()

    _, var_rows = build_variables_sheet(
        profile_rows, namemap_df, {}, cfg, total_rows=1010
    )
    assert var_rows
    summary = var_rows[0][9]
    assert summary == "fewer than 11 values", f"Got {summary!r}"


def test_num_text_no_cell_equals_n_plus_n_unparsed_when_unparsed_suppressed():
    """When N_UNPARSED is primary-suppressed, no VARIABLES cell must equal N + N_UNPARSED.

    E.g. N=1460, N_UNPARSED=<11 — cell must never be 1460 + any number 1-10.
    n_nonmiss must be 1460, not a sum that reveals N_UNPARSED.
    """
    profile_rows = _make_rows(
        {"variable": "PCNR_Z", "var_class": "NUM_TEXT", "stat": "N", "n": "1460", "small_cell_flag": "N"},
        {"variable": "PCNR_Z", "var_class": "NUM_TEXT", "stat": "N_UNPARSED", "n": "<11", "small_cell_flag": "Y"},
        {"variable": "PCNR_Z", "var_class": "NUM_TEXT", "stat": "N_MISSING", "n": "suppressed", "small_cell_flag": "N"},
    )
    namemap_df = pd.DataFrame(columns=["s_code", "pcnr_name", "original_header", "status"])
    cfg = _cfg()

    _, var_rows = build_variables_sheet(
        profile_rows, namemap_df, {}, cfg, total_rows=1470
    )
    assert var_rows
    row = var_rows[0]
    # Any integer cell 1461-1469 would reveal that N_UNPARSED is 1-9
    for cell in row:
        if isinstance(cell, int) and 1461 <= cell <= 1469:
            raise AssertionError(
                f"Cell value {cell} equals N + N_UNPARSED, revealing suppressed N_UNPARSED"
            )


# ---------------------------------------------------------------------------
# KEY sheet descriptions
# ---------------------------------------------------------------------------

def test_key_rows_contain_num_text_descriptions():
    """KEY rows must describe the current NUM_TEXT n_nonmiss and summary behaviour."""
    rows = _key_rows()
    all_text = " ".join(
        str(cell) for row in rows for cell in row if cell
    )

    assert "N + N_UNPARSED when both are visible; N only when N_UNPARSED is suppressed" in all_text, (
        "VARIABLES.n_nonmiss KEY description missing NUM_TEXT suppression rule"
    )
    assert "; <11 non-numeric entries' when that count is suppressed" in all_text, (
        "VARIABLES.summary KEY description missing NUM_TEXT non-numeric entries note"
    )

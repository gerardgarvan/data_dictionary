"""
Core logic for the name-map QC workbook (D-09, D-22).

Reads:   exchange/name_map.csv  (and other exchange CSVs)
Writes:  qc/name_map_review.xlsx
Sheets:  KEY | NAME_MAP | VARIABLES | NAME_MAP_COLLISIONS | NEAR_MATCHES | ABBREVIATIONS
"""

from __future__ import annotations

import csv
import logging
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook

from pdd.xlsx_style import style_sheet, write_table

EXPECTED_ROWS = 4795
SHEET_ORDER = [
    "KEY",
    "NAME_MAP",
    "VARIABLES",
    "NAME_MAP_COLLISIONS",
    "NEAR_MATCHES",
    "ABBREVIATIONS",
]

# Column names that should be written as integers in the workbook.
_NUMERIC_INT_COLS: frozenset[str] = frozenset([
    "col_order",
    "col_order_a",
    "col_order_b",
    "length",
    "max_length",
    "n_sources",
    "n_files_a",
    "n_files_b",
    "group_id",
    "group_id_a",
    "group_id_b",
    "edit_dist",
    "suffix_n",
    "trunc_flag",
])


def build_group_meta(
    name_map_rows: list[dict[str, str]],
) -> dict[str, dict[str, str]]:
    """For each group_id: representative original_header and sorted s_codes string.

    Representative = row with lowest s_code, then lowest col_order (numeric).
    """
    groups: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in name_map_rows:
        gid = row.get("group_id", "").strip()
        if gid:
            groups[gid].append(row)

    result: dict[str, dict[str, str]] = {}
    for gid, rows in groups.items():
        def _sort_key(r: dict[str, str]) -> tuple[str, int]:
            try:
                co = int(r.get("col_order", "0") or "0")
            except ValueError:
                co = 0
            return (r.get("s_code", ""), co)

        sorted_rows = sorted(rows, key=_sort_key)
        rep_header = sorted_rows[0].get("original_header", "")
        s_codes = sorted({r.get("s_code", "") for r in rows if r.get("s_code", "")})
        result[gid] = {"original_header": rep_header, "s_codes": ", ".join(s_codes)}
    return result


def _int_col_indices(headers: list[str]) -> frozenset[int]:
    return frozenset(i for i, h in enumerate(headers) if h in _NUMERIC_INT_COLS)


def read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    """Return (fieldnames, rows).  fieldnames is non-empty even for header-only files."""
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        rows = list(reader)
        headers = list(reader.fieldnames or [])
    return headers, rows


def assert_input(
    headers: list[str],
    name_map_rows: list[dict[str, str]],
    expected_rows: int = EXPECTED_ROWS,
) -> None:
    """Raise ValueError if required columns missing, row count wrong, or pcnr_name/group_id rules violated.

    Rules:
    - headers must contain 'pcnr_name' and 'group_id'.
    - len(name_map_rows) must equal expected_rows.
    - Any non-blank pcnr_name (case-insensitive) must map to exactly one
      distinct group_id. Blank pcnr_name values are unconstrained.
    """
    if "pcnr_name" not in headers:
        raise ValueError("name_map.csv missing required column: pcnr_name")
    if "group_id" not in headers:
        raise ValueError("name_map.csv missing required column: group_id")
    if len(name_map_rows) != expected_rows:
        raise ValueError(
            f"row count {len(name_map_rows)} != expected {expected_rows}"
        )
    # Build map: pcnr_name (lower) -> set of group_ids
    name_to_groups: dict[str, set[str]] = {}
    for r in name_map_rows:
        name = r.get("pcnr_name", "").strip()
        if not name:
            continue
        key = name.lower()
        group = r.get("group_id", "")
        if key not in name_to_groups:
            name_to_groups[key] = set()
        name_to_groups[key].add(group)
    for name_lower, groups in name_to_groups.items():
        if len(groups) > 1:
            raise ValueError(
                f"pcnr_name '{name_lower}' maps to multiple group_ids: {sorted(groups)}"
            )
    logging.info("Input asserts passed: %d rows", len(name_map_rows))


DECISIONS_TEMPLATE_COLS: list[str] = [
    "s_code_a", "header_a", "col_order_a",
    "s_code_b", "header_b", "col_order_b",
    "decision", "note", "new_name", "decided_by", "date",
]


def make_decisions_template(
    near_match_rows: list[dict[str, str]],
) -> list[dict[str, str]]:
    """One template row per near-match pair; decision columns are blank."""
    return [
        {col: row.get(col, "") for col in DECISIONS_TEMPLATE_COLS}
        for row in near_match_rows
    ]


def _key_rows() -> list[list[str]]:
    """Documentation rows for the KEY sheet."""
    return [
        ["Rule", "Item", "Description"],
        # D-02 cleaning steps
        ["D-02", "1", "Trim leading/trailing whitespace from the original header."],
        ["D-02", "2",
         "Replace every run of characters that are not ASCII letters or digits "
         "with a single underscore (non-ASCII treated as non-keeping)."],
        ["D-02", "3",
         "Strip leading and trailing underscores. "
         "A header that cleans to empty causes %fail_out."],
        ["D-02", "4", "If the result starts with a digit, prefix N_."],
        ["D-02", "5", "Keep original letter case."],
        ["D-02", "6",
         "Apply abbreviation table (D-03) only if the result would exceed 27 chars "
         "after PCNR_ prefix. Applies to all files."],
        ["D-02", "7", "Prefix PCNR_."],
        # D-03 abbreviation/elide rules
        ["D-03", "Abbrev",
         "Apply abbreviations from the ABBREVIATIONS sheet (docs/name_abbrev.csv) "
         "only when the cleaned name exceeds 27 chars. Apply longest abbreviation first. "
         "Token boundaries: start/end of string, underscore, digit, or lower-to-upper "
         "case change. Whole tokens only — never split mid-token."],
        ["D-03", "Elide",
         "If still > 27 chars after all abbreviations: keep first N and last M chars "
         "separated by __ (double underscore) to preserve COM/COPY prefix and suffix."],
        # Variable statuses
        ["Status", "PCNR",
         "Variable included in the PCNR_ namespace — appears in the final data dictionary."],
        ["Status", "KEY",
         "Linkage/identifier variable retained as-is; not renamed to PCNR_."],
        ["Status", "LINKAGE_ONLY",
         "Variable used for record linkage only; excluded from analytic output."],
        ["Status", "EXCLUDED_CLK",
         "Variable excluded because it is a clock-time or date field raising "
         "re-identification risk."],
        ["Status", "DUPLICATE_DROPPED",
         "Variable is an exact duplicate of another in the same file; "
         "dropped before naming."],
        # Sheet descriptions
        ["Sheet", "NAME_MAP",
         "One row per variable per source file. Shows the cleaned name, status, "
         "and group_id assigned by the SAS pipeline."],
        ["Sheet", "VARIABLES",
         "One row per unique PCNR_ variable across all source files. Shows type, "
         "max length, how many sources contribute it, plus original_header and s_codes."],
        ["Sheet", "NAME_MAP_COLLISIONS",
         "Groups of variables whose cleaned names collide (same PCNR_ base name). "
         "Includes original_header and s_codes for the group."],
        ["Sheet", "NEAR_MATCHES",
         "Pairs of PCNR_ names across files that are similar but not identical — "
         "candidates for manual harmonisation."],
        ["Sheet", "ABBREVIATIONS",
         "Abbreviation pairs applied by D-03 (docs/name_abbrev.csv). "
         "long = original token, short = replacement."],
        # group_id
        ["group_id", "Meaning",
         "Each distinct PCNR_ variable has one group_id. Rows in NAME_MAP from "
         "different source files that share a group_id are overlaps — the same "
         "real-world variable appearing in multiple files."],
        # How to record a decision
        ["Decision", "Columns",
         "docs/name_map_decisions.csv: s_code_a, header_a, col_order_a, "
         "s_code_b, header_b, col_order_b, decision, note, new_name, decided_by, date."],
        ["Decision", "SAME",
         "The two near-match variables are the same construct — keep existing group_id; "
         "no rename needed."],
        ["Decision", "DIFFERENT",
         "The two near-match variables are distinct constructs — keep separate group_ids."],
        ["Decision", "RENAME",
         "One variable should be renamed; put the chosen PCNR_ name in new_name."],
        # Derived columns added by Python post-processing
        ["Column", "original_header",
         "Representative original header text for this group_id — the row with the "
         "lowest s_code, then lowest col_order (numeric)."],
        ["Column", "s_codes",
         "Comma-separated distinct source codes (s_code) contributing to this "
         "group_id, sorted alphabetically."],
    ]


def write_workbook(
    sheets: dict[str, tuple[list[str], list[list]]],
    out_path: Path,
) -> None:
    """Write the name-map review workbook.

    Parameters
    ----------
    sheets:
        Mapping of sheet name to (headers, data_rows). KEY sheet is always
        written as the hard-coded D-02 cleaning-rule text; KEY entry in
        ``sheets`` is ignored if present.
    out_path:
        Destination path for the xlsx file. Parent directories are created.
    """
    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]

    for sheet_name in SHEET_ORDER:
        ws = wb.create_sheet(sheet_name)
        if sheet_name == "KEY":
            key_data = _key_rows()
            headers = key_data[0]
            data = key_data[1:]
            write_table(ws, headers, data)
        elif sheet_name in sheets:
            headers, data_rows = sheets[sheet_name]
            int_cols = _int_col_indices(headers)
            write_table(ws, headers, data_rows, int_cols=int_cols)
        else:
            ws.append(["(no data)"])
        style_sheet(ws)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    logging.info("Saved %s", out_path)


def verify_workbook(path: Path, expected_data_rows: dict[str, int]) -> None:
    """Re-open and confirm expected sheets, order, and row counts.

    Raises ValueError on any mismatch.
    """
    wb = load_workbook(path, read_only=True)
    try:
        actual_names = wb.sheetnames
        if actual_names != SHEET_ORDER:
            raise ValueError(
                f"sheet order mismatch: {actual_names} != {SHEET_ORDER}"
            )
        for sheet_name, exp in expected_data_rows.items():
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"sheet '{sheet_name}' not found in workbook")
            ws = wb[sheet_name]
            actual = ws.max_row - 1  # subtract header
            if actual != exp:
                raise ValueError(
                    f"sheet {sheet_name}: {actual} data rows != expected {exp}"
                )
    finally:
        wb.close()
    logging.info("Workbook verified: %s", path)

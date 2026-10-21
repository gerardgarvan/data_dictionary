"""
Core logic for the s06 clock-classification QC workbook (D-21, D-22).

Reads:   exchange/s06_class_nn.csv
Writes:  qc/s06_clock_classification.xlsx
Sheets:  KEY | SUMMARY | CLOCK | DIGITAL_OTHER | PCNR | ALL_COLUMNS
"""

from __future__ import annotations

import csv
import logging
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook

from pdd.xlsx_style import (
    ACCENT_FILL,
    ACCENT_FONT,
    apply_header_row,
    autofit_columns,
    freeze_and_filter,
    write_table,
)

EXPECTED_ROWS = 3987
EXPECTED_COUNTS: dict[str, int] = {
    "CLOCK": 2747,
    "DIGITAL_OTHER": 809,
    "PCNR": 431,
    "UNSURE": 0,
}

SHEET_ORDER = ["KEY", "SUMMARY", "CLOCK", "DIGITAL_OTHER", "PCNR", "ALL_COLUMNS"]

_COUNT_FMT = "#,##0"
_PCT_FMT = "0.0"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def assert_input(
    rows: list[dict[str, str]],
    expected_rows: int = EXPECTED_ROWS,
    expected_counts: dict[str, int] | None = None,
) -> None:
    """Raise ValueError if row count or class counts do not match expectations."""
    if expected_counts is None:
        expected_counts = EXPECTED_COUNTS
    if len(rows) != expected_rows:
        raise ValueError(f"row count {len(rows)} != expected {expected_rows}")
    counts = Counter(r["class"] for r in rows)
    for cls, exp in expected_counts.items():
        actual = counts.get(cls, 0)
        if actual != exp:
            raise ValueError(f"class={cls} count={actual} != expected {exp}")
    logging.info("Input asserts passed: %d rows, %s", len(rows), dict(counts))


def _key_rows() -> list[list[str]]:
    return [
        ["Output column", "sas_name", "SAS variable name in pdd.src_s06 (max 32 chars; from pdd.import_spec)"],
        ["Output column", "orig_col_name", "Original column header from the source file"],
        ["Output column", "token", "Text after COM/COPY prefix; LINUS for LINUS-prefix columns; blank otherwise"],
        ["Output column", "class", "Classification: CLOCK | DIGITAL_OTHER | PCNR | UNSURE"],
        ["Output column", "rule", "Text description of the matching classification rule"],
        ["Output column", "pct_nonnull", "% rows in pdd.src_s06 with a non-missing value (0-100)"],
        ["Class value", "CLOCK", "Clock-drawing features: COM/COPY + clock token OR LINUS DCTclock. Excluded from PCNR_. n=2,747 (D-21)."],
        ["Class value", "DIGITAL_OTHER", "Other digital cognitive battery tests: COM/COPY + non-clock digital token. Stay in PCNR_. n=809."],
        ["Class value", "PCNR", "All other source columns. Always PCNR_. n=431. Includes clock admin flags and Mini-Cog scores."],
        ["Class value", "UNSURE", "COM/COPY prefix with unrecognized token. Must be 0 after D-21 review."],
    ]


def _summary_tables(
    rows: list[dict[str, str]],
) -> tuple[list[list], list[list]]:
    class_counts = Counter(r["class"] for r in rows)
    # counts are int, not str
    class_rows: list[list] = [[cls, class_counts.get(cls, 0)] for cls in ("CLOCK", "DIGITAL_OTHER", "PCNR", "UNSURE")]

    rule_class: dict[tuple[str, str], int] = {}
    for r in rows:
        key = (r["rule"], r["class"])
        rule_class[key] = rule_class.get(key, 0) + 1
    rule_rows: list[list] = [[rule, cls, n] for (rule, cls), n in sorted(rule_class.items())]

    return class_rows, rule_rows


def _subset_data(subset: list[dict[str, str]]) -> list[list]:
    return [
        [
            r["sas_name"],
            r["orig_col_name"],
            r["token"],
            r["rule"],
            float(r["pct_nonnull"]) if r["pct_nonnull"] else None,
        ]
        for r in subset
    ]


def _accent_column(ws, col_idx: int) -> None:
    """Orange accent on data cells (row 2+) in the given 1-based column."""
    for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
        for cell in row:
            cell.fill = ACCENT_FILL
            cell.font = ACCENT_FONT


def _write_summary(ws, class_rows: list[list], rule_rows: list[list]) -> None:
    """Two-table SUMMARY sheet: class counts (with Total row) then rule breakdown."""
    # --- Class count table ---
    ws.append(["Class", "Column count"])
    apply_header_row(ws)
    for row in class_rows:
        ws.append(row)
    # Total row styled with UF orange
    total = sum(r[1] for r in class_rows)
    total_row_idx = ws.max_row + 1
    ws.append(["Total", total])
    for cell in ws[total_row_idx]:
        cell.fill = ACCENT_FILL
        cell.font = ACCENT_FONT
    # Apply count format to column 2 (rows 2 onward)
    for row_cells in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row_cells:
            if cell.value is not None:
                cell.number_format = _COUNT_FMT
    # Orange on Class column data cells (col 1, rows 2 to before Total)
    for row_cells in ws.iter_rows(min_row=2, min_col=1, max_col=1, max_row=total_row_idx - 1):
        for cell in row_cells:
            cell.fill = ACCENT_FILL
            cell.font = ACCENT_FONT
    # --- Rule breakdown table ---
    ws.append([])
    ws.append(["Rule", "Class", "Column count"])
    for row in rule_rows:
        ws.append(row)
    # Apply count format to column 3 in rule section
    rule_header_row = total_row_idx + 2
    for row_cells in ws.iter_rows(min_row=rule_header_row + 1, min_col=3, max_col=3):
        for cell in row_cells:
            if cell.value is not None:
                cell.number_format = _COUNT_FMT
    freeze_and_filter(ws)
    autofit_columns(ws)


def write_workbook(rows: list[dict[str, str]], out_path: Path) -> None:
    """Write the classification workbook; does not validate row counts."""
    clock = [r for r in rows if r["class"] == "CLOCK"]
    do = [r for r in rows if r["class"] == "DIGITAL_OTHER"]
    pcnr = [r for r in rows if r["class"] == "PCNR"]
    class_rows, rule_rows = _summary_tables(rows)

    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]

    # KEY
    ws_key = wb.create_sheet("KEY")
    write_table(ws_key, ["Section", "Item", "Description"], _key_rows())

    # SUMMARY: two tables in one sheet (special handling)
    ws_sum = wb.create_sheet("SUMMARY")
    _write_summary(ws_sum, class_rows, rule_rows)

    sub_headers = ["SAS Name", "Original Column", "Token", "Rule", "% Non-null"]
    _pct_col = frozenset({4})
    _pct_fmt = {4: _PCT_FMT}

    ws_clock = wb.create_sheet("CLOCK")
    write_table(ws_clock, sub_headers, _subset_data(clock), float_cols=_pct_col, formats=_pct_fmt)

    ws_do = wb.create_sheet("DIGITAL_OTHER")
    write_table(ws_do, sub_headers, _subset_data(do), float_cols=_pct_col, formats=_pct_fmt)

    ws_pcnr = wb.create_sheet("PCNR")
    write_table(ws_pcnr, sub_headers, _subset_data(pcnr), float_cols=_pct_col, formats=_pct_fmt)

    ws_all = wb.create_sheet("ALL_COLUMNS")
    all_data = [
        [r["sas_name"], r["orig_col_name"], r["token"], r["class"], r["rule"],
         float(r["pct_nonnull"]) if r["pct_nonnull"] else None]
        for r in rows
    ]
    write_table(
        ws_all,
        ["SAS Name", "Original Column", "Token", "Class", "Rule", "% Non-null"],
        all_data,
        float_cols=frozenset({5}),
        formats={5: _PCT_FMT},
    )
    _accent_column(ws_all, 4)  # Class column orange

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    logging.info("Saved %s", out_path)


def verify_workbook(path: Path, expected_data_rows: dict[str, int]) -> None:
    """Re-open and confirm expected sheets and row counts. Raises ValueError on mismatch."""
    wb = load_workbook(path, read_only=True)
    try:
        missing = [s for s in SHEET_ORDER if s not in wb.sheetnames]
        if missing:
            raise ValueError(f"missing sheets: {missing}")
        for sheet_name, exp in expected_data_rows.items():
            ws = wb[sheet_name]
            actual = ws.max_row - 1  # subtract header
            if actual != exp:
                raise ValueError(f"sheet {sheet_name}: {actual} data rows != expected {exp}")
    finally:
        wb.close()
    logging.info("Workbook verified: %s", path)

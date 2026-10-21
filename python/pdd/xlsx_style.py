from __future__ import annotations

from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill("solid", fgColor="0021A5")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ACCENT_FILL = PatternFill("solid", fgColor="FA4616")
ACCENT_FONT = Font(color="FFFFFF")


def apply_header_row(ws: Worksheet) -> None:
    """Apply UF blue fill + white bold text to every cell in row 1."""
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT


def freeze_and_filter(ws: Worksheet) -> None:
    """Freeze row 1 and set autofilter on all data columns."""
    ws.freeze_panes = "A2"
    if ws.max_row > 0 and ws.max_column > 0:
        ws.auto_filter.ref = ws.dimensions


def autofit_columns(ws: Worksheet, min_width: int = 8, max_width: int = 60) -> None:
    """Set column width from max content length."""
    for col in ws.columns:
        width = max(
            (len(str(cell.value)) if cell.value is not None else 0 for cell in col),
            default=min_width,
        )
        ws.column_dimensions[col[0].column_letter].width = min(
            max(width + 2, min_width), max_width
        )


def style_sheet(ws: Worksheet) -> None:
    """Apply header row, freeze, autofilter, and autofit in one call."""
    apply_header_row(ws)
    freeze_and_filter(ws)
    autofit_columns(ws)
